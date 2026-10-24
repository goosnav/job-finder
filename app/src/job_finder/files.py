"""Bounded document imports and spreadsheet-safe exports."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

from .config import AppError
from .quality import candidate_signature

MAX_FILE = 8 * 1024 * 1024
MAX_ROWS = 10000


def check_file(raw: bytes):
    if not raw:
        raise AppError("FILE-EMPTY", "Choose a file with content.")
    if len(raw) > MAX_FILE:
        raise AppError("FILE-SIZE", "Files must be 8 MB or smaller.", 413)


def check_zip(raw: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            items = archive.infolist()
            if len(items) > 2000 or sum(i.file_size for i in items) > 40 * 1024 * 1024:
                raise AppError("FILE-SIZE", "The expanded document is too large.")
            if any(i.flag_bits & 1 for i in items):
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted document.")
    except zipfile.BadZipFile as exc:
        raise AppError("FILE-FORMAT", "This document is not a valid Office file.") from exc


def document_text(name: str, raw: bytes) -> str:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv"}:
            text = raw.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise AppError("FILE-ENCRYPTED", "Use an unencrypted PDF.")
            if len(reader.pages) > 80:
                raise AppError("FILE-SIZE", "Use a PDF with 80 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            check_zip(raw)
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                xml = archive.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            text = "\n".join(
                "".join(p.itertext())
                for p in root.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
            )
        else:
            raise AppError("FILE-TYPE", "Use TXT, Markdown, PDF, DOCX, or CSV for your résumé or supporting material.")
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "The document could not be read. Try a plain-text export."
        ) from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise AppError("FILE-NO-TEXT", "No text found. Scanned PDFs need OCR before upload.")
    if len(text) > 60000:
        raise AppError("FILE-SIZE", "Use a shorter document (up to 60,000 text characters).")
    return text


def header(value) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def read_rows(name: str, raw: bytes) -> list[dict]:
    check_file(raw)
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".csv":
            text = raw.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text, newline=""), dialect)
            rows = list(_bounded(reader))
        elif suffix == ".xlsx":
            check_zip(raw)
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                rows = list(_bounded(book.active.iter_rows(values_only=True)))
            finally:
                book.close()
        else:
            raise AppError("FILE-TYPE", "Import a CSV or XLSX file.")
        if len(rows) < 2:
            raise AppError("FILE-ROWS", "Include a header and at least one data row.")
        keys = [header(k) for k in rows[0]]
        if any(not key for key in keys) or len(keys) != len(set(keys)):
            raise AppError("FILE-HEADERS", "Use unique, nonempty column names.")
        return [
            {
                key: str(value).strip() if value is not None else ""
                for key, value in zip(keys, row, strict=False)
                if key
            }
            for row in rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "FILE-READ", "Could not read this spreadsheet. Try a UTF-8 CSV export."
        ) from exc


def _bounded(rows):
    for index, row in enumerate(rows):
        if index > MAX_ROWS or len(row) > 150:
            raise AppError("FILE-SIZE", "Imports support up to 10,000 rows and 150 columns.")
        if any(len(str(value or "")) > 60000 for value in row):
            raise AppError("FILE-SIZE", "A cell is too large.")
        yield row


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
            return "'" + value
    return value


BASE_COLUMNS = [
    "rank", "fit_score", "confidence", "eligibility", "next_action", "company", "title", "kind",
    "company_url", "listing_url", "apply_url", "careers_url", "requisition_id", "location",
    "work_mode", "remote_eligibility", "employment_type", "seniority", "department", "industry",
    "salary_min", "salary_max", "salary_currency", "salary_period", "salary_source",
    "sponsorship", "authorization_requirement", "clearance_requirement", "posted_at",
    "date_basis", "source_updated_at", "valid_through", "listing_status", "listing_source",
    "last_checked_at", "source_check", "description", "requirements", "score_breakdown",
    "blockers", "questions", "skill_gaps", "research_summary", "interview_hooks", "application_tips",
    "signals", "candidate_claims", "outreach_angle", "risks", "contact_name", "contact_role",
    "contact_source", "email", "email_status", "email_source", "email_checked_at",
    "verification_source", "verification_at", "linkedin_url", "research_status", "status",
    "ready_to_apply", "apply_blockers", "ready_to_contact", "export_blockers", "review_notes",
    "application_status", "activity_at", "next_follow_up", "prior_activity",
]
EXPORT_COLUMNS = BASE_COLUMNS + [
    f"email_{step}_{field}" for step in (1, 2, 3)
    for field in ("subject", "body", "delay_days", "source_urls", "candidate_claims")
] + ["candidate_name", "reply_email", "signature", "source_urls", "evidence", "created_at", "updated_at", "opportunity_id", "search_id"]


def export_rows(search: dict, opportunities: list[dict]) -> list[dict]:
    result = []
    signature = candidate_signature(search["candidate"])
    for opportunity in opportunities:
        row = {key: opportunity.get(key, "") for key in BASE_COLUMNS}
        for key in ("blockers", "questions", "apply_blockers", "export_blockers"):
            row[key] = "; ".join(opportunity.get(key, []))
        for draft in opportunity.get("drafts", []):
            for field in ("subject", "body", "delay_days", "source_urls", "candidate_claims"):
                value = draft.get(field, "")
                if field == "body" and signature:
                    value += "\n\n" + signature
                row[f"email_{draft['step']}_{field}"] = value
        row.update(candidate_name=search["candidate"]["name"], reply_email=search["candidate"]["email"], signature=signature)
        for key in ("source_urls", "evidence", "created_at", "updated_at", "search_id"):
            row[key] = opportunity.get(key, "")
        row["opportunity_id"] = opportunity["id"]
        result.append(row)
    return result


def spreadsheet(search: dict, opportunities: list[dict], format: str) -> bytes:
    rows = export_rows(search, opportunities)
    if format == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=EXPORT_COLUMNS)
        writer.writeheader()
        writer.writerows({key: safe_cell(row.get(key, "")) for key in EXPORT_COLUMNS} for row in rows)
        return buffer.getvalue().encode("utf-8-sig")
    if format != "xlsx":
        raise AppError("EXPORT-FORMAT", "Choose CSV or XLSX.")
    book = Workbook()
    summary = book.active
    summary.title = "Report"
    preferences = search["preferences"]
    for row in [
        ["JOB FINDER", "Goosnav Labs"],
        ["Search", search["name"]],
        ["Objective", search["objective"]],
        ["Target roles", "; ".join(preferences["roles"]) or "Not specified"],
        ["Locations", "; ".join(preferences["locations"]) or "Not specified"],
        ["Work modes", "; ".join(preferences["work_modes"]) or "Not specified"],
        ["Minimum annual salary", f"{preferences['salary_currency']} {preferences['salary_min']:,.0f}" if preferences["salary_min"] is not None else "Not specified"],
        ["Rows in this report", len(opportunities)],
        ["Data", "FICTIONAL DEMO" if search["demo"] else "Review source dates and unresolved requirements before applying."],
    ]:
        summary.append([safe_cell(value) for value in row])
    summary.append([])
    summary.append(["Rank", "Role", "Company", "Location", "Fit / 100", "Listing status", "Progress", "Job link"])
    for opportunity in opportunities:
        summary.append([safe_cell(value) for value in (
            opportunity.get("rank", ""), opportunity.get("title", ""), opportunity["company"],
            opportunity.get("location", ""), opportunity.get("fit_score", ""),
            opportunity.get("listing_status", "unverified"), opportunity.get("application_status", "new"),
            opportunity.get("apply_url") or opportunity.get("listing_url") or opportunity.get("careers_url", ""),
        )])
    summary.freeze_panes = "D12"
    summary.auto_filter.ref = f"A11:H{max(11, summary.max_row)}"
    for column, width in {"A": 24, "B": 48, "C": 30, "D": 30, "E": 15, "F": 20, "G": 20, "H": 60}.items():
        summary.column_dimensions[column].width = width
    for row in summary:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row in {1, 11}:
                cell.fill = PatternFill("solid", fgColor="1746C7")
                cell.font = Font(color="FFFFFF", bold=True)
    for cell in list(summary.columns)[7][11:]:
        if str(cell.value or "").startswith(("https://", "http://")):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
    sheet = book.create_sheet("Opportunities")
    sheet.append(EXPORT_COLUMNS)
    for row in rows:
        sheet.append([safe_cell(row.get(key, "")) for key in EXPORT_COLUMNS])
    sheet.freeze_panes = "H2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="172320")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[cell.column_letter].width = 25
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sources = book.create_sheet("Sources")
    sources.append(["Opportunity ID", "Company", "Role", "Field", "Source URL", "Excerpt", "Checked at"])
    for opportunity in opportunities:
        for evidence in opportunity.get("evidence", []):
            sources.append([safe_cell(x) for x in (opportunity["id"], opportunity["company"], opportunity["title"], evidence["field"], evidence["url"], evidence["quote"], evidence["checked_at"])])
    sources.freeze_panes = "A2"
    sources.auto_filter.ref = sources.dimensions
    for cell in sources[1]:
        cell.fill = PatternFill("solid", fgColor="172320")
        cell.font = Font(color="FFFFFF", bold=True)
        sources.column_dimensions[cell.column_letter].width = 30
    notes = book.create_sheet("Read me")
    for row in [
        ["Search", search["name"]],
        ["Data", "FICTIONAL DEMO — DO NOT APPLY OR CONTACT" if search["demo"] else "Research and drafts. No applications or messages have been sent."],
        ["Ranking", "Fit is a weighted evidence-backed match, not a probability of getting hired. Unknown facts earn no fit points. Confidence measures assessed evidence coverage, not verification accuracy."],
        ["Company leads", "Funding, expansion and other signals suggest possible needs. They do not prove a vacancy."],
        ["Apply", "Use ready_to_apply=yes for recently checked direct application links. Review unresolved requirements before submitting."],
        ["Outreach", "Use ready_to_contact=yes only after review. Email bodies already include the candidate signature. One recipient per outreach export."],
        ["Verification", "Published means a public address was sourced. MX-found confirms domain mail routing only, not mailbox deliverability."],
        ["Evidence", "Agent-reported facts require judgment. Review all résumé claims, requirements, source excerpts and posting dates."],
        ["Follow-up", "Delay days are measured from first contact. Stop after any reply, bounce or opt-out. Tracking is manual or imported; no automatic sending."],
        ["Privacy", "This export contains your supplied candidate name, reply email and selected résumé claims. Review before sharing."],
    ]:
        notes.append([safe_cell(x) for x in row])
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 110
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
