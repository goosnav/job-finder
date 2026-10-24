#!/usr/bin/env python3
"""Check production onboarding and reports against an isolated local workspace."""

import argparse
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import time
import urllib.request
from contextlib import contextmanager
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

RESUME = b'Fictional candidate for application testing only. Mechanical engineer with five years of CAD, prototype design, test fixtures, and Python data analysis.'


@contextmanager
def workspace(url):
    if url:
        yield url
        return
    root = Path(__file__).resolve().parents[1]
    with tempfile.TemporaryDirectory(prefix='job-finder-gui-') as data:
        env = dict(os.environ, JOB_FINDER_DATA=data, PYTHONPATH=str(root / 'app/src'))
        child = subprocess.Popen([sys.executable, '-u', '-c', 'from job_finder.start import serve; serve(open_browser=False)'], cwd=root, env=env, stdout=subprocess.PIPE, text=True)
        try:
            line = child.stdout.readline().strip()
            if not line.startswith('Job Finder: http://127.0.0.1:'):
                raise RuntimeError('The isolated GUI did not start.')
            url = line.removeprefix('Job Finder: ')
            for _ in range(50):
                try:
                    with urllib.request.urlopen(url + '/health/ready', timeout=1) as response:
                        if json.load(response)['ready']:
                            break
                except OSError:
                    time.sleep(.1)
            else:
                raise RuntimeError('The isolated GUI did not become ready.')
            yield url
        finally:
            if child.poll() is None:
                child.terminate()
                try:
                    child.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    child.kill()
                    child.wait(timeout=5)


def run_checks(args):
    args.output.mkdir(parents=True, exist_ok=True)
    failures = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(viewport={'width': 1440, 'height': 1000})
        page = context.new_page()
        page.on('pageerror', lambda error: failures.append(str(error)))
        page.on('dialog', lambda dialog: dialog.accept())
        page.on('response', lambda response: failures.append(f'{response.status} {response.url}')
                if response.status >= 400 and not (response.status in {400, 422} and any(path in response.url for path in ('/intake', '/documents'))) else None)
        page.goto(args.url)
        page.wait_for_load_state('networkidle')
        # Legacy demo data must never become the production starting screen.
        token = page.request.get(args.url + '/api/session').json()['token']
        sample = page.request.post(args.url + '/api/demo', headers={'X-Session-Token': token}).json()['search']
        page.evaluate('(id) => localStorage.setItem("job-finder-active-search",id)', sample['id'])
        page.reload()
        if page.locator('#active-search option').count():
            page.get_by_role('button', name='New search', exact=False).click()
        expect(page.locator('#intake-resume')).to_be_visible()
        assert not page.get_by_role('button', name='Load sample').count()
        assert not page.locator(f'#active-search option[value="{sample["id"]}"]').count()
        for width in (320, 390, 768, 1440):
            page.set_viewport_size({'width': width, 'height': 1000})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), width
            page.screenshot(path=str(args.output / ('intake-mobile.png' if width == 390 else 'intake-desktop.png')), full_page=True)
        # Invalid resume extraction must retain all inputs and create no partial search.
        page.get_by_label('Locations', exact=True).fill('Los Angeles, CA; San Diego, CA')
        page.get_by_label('Minimum annual salary', exact=True).fill('120000')
        page.get_by_label('hybrid', exact=True).check()
        page.locator('#intake-resume').set_input_files({'name': 'broken.pdf', 'mimeType': 'application/pdf', 'buffer': b'invalid pdf'})
        before = len(page.request.get(args.url + '/api/searches').json()['searches'])
        page.get_by_role('button', name='Save search', exact=True).click()
        expect(page.locator('#intake-error')).to_be_visible()
        expect(page.locator('#new-search-form button[type=submit]:disabled')).to_have_count(0)
        expect(page.get_by_label('Minimum annual salary', exact=True)).to_have_value('120000')
        assert len(page.request.get(args.url + '/api/searches').json()['searches']) == before
        # Resume-only input is valid; goals are optional.
        page.locator('#intake-resume').set_input_files({'name': 'synthetic-resume.txt', 'mimeType': 'text/plain', 'buffer': RESUME})
        # A pending intake must reject both a second click and requestSubmit().
        intake_requests = []
        def track_intake(route):
            intake_requests.append(route.request.url)
            time.sleep(.3)
            route.continue_()
        page.route('**/api/searches/intake', track_intake)
        page.evaluate('''(plan) => {
            const form = document.querySelector('#new-search-form');
            const buttons = form.querySelectorAll('button[type=submit]');
            buttons[plan ? 1 : 0].click();
            if (![...buttons].every(button => button.disabled)) throw Error('Intake buttons stayed enabled');
            buttons[plan ? 0 : 1].click();
            form.requestSubmit(buttons[plan ? 0 : 1]);
        }''', args.run_agent)
        if args.run_agent:
            expect(page.locator('#agent-consent-dialog')).to_be_visible()
            page.get_by_role('button', name='Allow & run', exact=True).click()
            expect(page.locator('.search-plan')).to_be_visible(timeout=180000)
            expect(page.locator('.plan-roles li').first).to_be_visible()
            expect(page.locator('.plan-constraints')).to_contain_text('USD 120,000')
            page.screenshot(path=str(args.output / 'search-plan.png'), full_page=True)
            page.get_by_role('button', name='Inputs', exact=True).click()
        else:
            expect(page.locator('#search-form')).to_be_visible()
        assert len(intake_requests) == 1
        assert len(page.request.get(args.url + '/api/searches').json()['searches']) == before + 1
        page.unroute('**/api/searches/intake', track_intake)
        expect(page.locator('.document-item')).to_have_count(1)
        expect(page.get_by_label('Locations', exact=True)).to_have_value('Los Angeles, CA; San Diego, CA')
        expect(page.get_by_label('Minimum annual salary', exact=True)).to_have_value('120000')
        # Mixed upload failure must reveal the accepted document.
        page.locator('#documents-input').set_input_files([
            {'name': 'extra.txt', 'mimeType': 'text/plain', 'buffer': b'Fictional portfolio: fixture design.'},
            {'name': 'unreadable.pdf', 'mimeType': 'application/pdf', 'buffer': b'invalid pdf'},
        ])
        expect(page.locator('.document-item')).to_have_count(2)
        expect(page.get_by_text('unreadable.pdf:', exact=False)).to_be_visible()
        page.get_by_role('button', name='Remove extra.txt', exact=True).click()
        expect(page.locator('.document-item')).to_have_count(1)
        page.screenshot(path=str(args.output / 'search.png'), full_page=True)
        page.locator('.agent-settings').evaluate('(el)=>el.open=true')
        page.locator('.route-grid').screenshot(path=str(args.output / 'agent-routes.png'))
        page.locator('.nav-tab[data-view=opportunities]').click()
        expect(page.get_by_role('button', name='CSV', exact=True)).to_be_disabled()
        page.get_by_role('button', name='Import', exact=True).click()
        page.locator('#import-file').set_input_files({'name': 'jobs.csv', 'mimeType': 'text/csv', 'buffer': b'company,title,listing_url,apply_url,location\nFictional QA Employer,Mechanical Engineer,https://jobs.example.invalid/role-1,https://jobs.example.invalid/apply-1,"Los Angeles, CA"\nFictional QA Employer,Test Engineer,https://jobs.example.invalid/role-2,https://jobs.example.invalid/apply-2,"San Diego, CA"\n'})
        page.locator('#import-form').get_by_role('button', name='Import', exact=True).click()
        expect(page.locator('tr[data-opportunity]')).to_have_count(2)
        assert page.locator('tr[data-opportunity]').first.bounding_box()['y'] < 300
        # Keyboard selection must retain focus after the table updates.
        page.locator('[data-select]').first.focus()
        page.keyboard.press('Space')
        expect(page.locator('[data-select]').first).to_be_focused()
        page.keyboard.press('Tab')
        assert page.locator('body').evaluate('(el)=>document.activeElement!==el')
        page.get_by_role('button', name='Clear', exact=True).click()
        page.get_by_role('button', name='Save to shortlist', exact=True).first.click()
        expect(page.get_by_role('button', name='Remove from shortlist', exact=True)).to_have_count(1)
        page.locator('.nav-tab[data-view=shortlist]').click()
        expect(page.locator('tr[data-opportunity]')).to_have_count(1)
        with page.expect_download() as downloaded:
            page.get_by_role('button', name='CSV', exact=True).click()
        assert len(list(csv.DictReader(io.StringIO(Path(downloaded.value.path()).read_text('utf-8-sig'))))) == 1
        page.locator('.nav-tab[data-view=opportunities]').click()
        page.locator('[data-action=open-opportunity]').first.click()
        page.get_by_label('Your notes', exact=True).fill('QA: check the role requirements.')
        page.get_by_role('button', name='Save edits', exact=True).click()
        expect(page.locator('#opportunity-save-indicator')).to_have_text('Saved locally')
        page.screenshot(path=str(args.output / 'opportunity-review.png'), full_page=True)
        page.get_by_role('button', name='Close opportunity').click()
        page.screenshot(path=str(args.output / 'opportunities.png'), full_page=True)
        with page.expect_download() as downloaded:
            page.get_by_role('button', name='CSV', exact=True).click()
        rows = list(csv.DictReader(io.StringIO(Path(downloaded.value.path()).read_text('utf-8-sig'))))
        assert len(rows) == 2 and rows[0]['apply_url'] and rows[0]['review_notes'].startswith('QA:')
        with page.expect_download() as downloaded:
            page.get_by_role('button', name='Excel', exact=True).click()
        book = load_workbook(io.BytesIO(Path(downloaded.value.path()).read_bytes()))
        assert book.sheetnames == ['Report', 'Opportunities', 'Sources', 'Read me']
        assert book['Report']['H12'].hyperlink
        job_link = page.get_by_role('link', name='Open job', exact=False).first
        context.route('https://jobs.example.invalid/**', lambda route: route.fulfill(body='Fictional application page for the GUI test.'))
        with page.expect_popup() as popup:
            job_link.click()
        expect(popup.value.locator('body')).to_contain_text('Fictional application page')
        popup.value.close()
        for width in (320, 390, 768):
            page.set_viewport_size({'width': width, 'height': 900})
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), width
            if width == 390:
                page.screenshot(path=str(args.output / 'mobile.png'), full_page=True)
        page.reload()
        expect(page.locator('tr[data-opportunity]')).to_have_count(2)
        # A goals-only search uses the same intake and persistence contract.
        page.get_by_role('button', name='New search', exact=False).click()
        page.get_by_label('Target roles', exact=True).fill('Product Engineer')
        page.get_by_label('Description', exact=True).fill('Hands-on product development; no management.')
        page.get_by_label('Locations', exact=True).fill('Remote US')
        page.get_by_label('Minimum annual salary', exact=True).fill('100000')
        page.get_by_role('button', name='Save search', exact=True).click()
        expect(page.locator('#search-form')).to_be_visible()
        expect(page.get_by_label('Target roles', exact=True)).to_have_value('Product Engineer')
        assert not failures, failures
        browser.close()
    print('PASS: production startup; resume-only and goals-only intake; extraction failure; partial uploads; plan' + (' (live Codex)' if args.run_agent else ' (not run)') + '; preferences; keyboard; shortlist; job access; CSV/XLSX; responsive; reopen. Zero unexpected JS/HTTP errors.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', help='Existing loopback GUI with isolated test data; default starts a temporary workspace')
    parser.add_argument('--output', type=Path, default=Path('dev/qa'))
    parser.add_argument('--run-agent', action='store_true', help='Run one real profile call with the synthetic resume')
    args = parser.parse_args()
    if args.url and not args.url.startswith('http://127.0.0.1:'):
        parser.error('Use an isolated loopback application.')
    with workspace(args.url) as url:
        args.url = url
        run_checks(args)


if __name__ == '__main__':
    main()
