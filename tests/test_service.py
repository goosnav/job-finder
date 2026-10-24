import csv
import io
from datetime import UTC, datetime, timedelta

import pytest
from job_finder import demo
from job_finder.config import AppError
from job_finder.models import OpportunitySeed, now, public_url, route_for
from job_finder.pipeline import Jobs
from job_finder.prompts import candidate_excerpt, queries
from job_finder.ranking import assess, candidate_claims_supported, deadline_passed, identity_of
from job_finder.service import Service
from openpyxl import load_workbook


@pytest.fixture
def service(tmp_path):
    return Service(tmp_path)


@pytest.fixture
def prepared(service):
    search = service.create_search({"name": "Fixture", "objective": "Operations engineering", "candidate": demo.candidate(), "preferences": {"roles": ["Operations Engineer"], "locations": ["Los Angeles, CA"], "salary_min": 100000}})
    seed = OpportunitySeed.model_validate(demo.discovery_result(1, [])['opportunities'][0])
    opportunity, _ = service.upsert_seed(search['id'], seed)
    opportunity.update(demo.research_result(opportunity))
    opportunity.update(research_status='complete', status='drafted', source_check='matched', last_checked_at=now(), drafts=demo.sequence_result(opportunity)['drafts'])
    opportunity.update(assess(search, opportunity, service.candidate_text(search['id'])))
    service.store.save('opportunities', opportunity)
    return search, opportunity


def ready(service, opportunity):
    return service.update_opportunity(opportunity['id'], {'email_status': 'published', 'status': 'approved'})


def test_demo_is_durable_ranked_and_cannot_enter_action_exports(service):
    search = service.create_demo()
    run = Jobs(service).create(search['id'], background=False)
    assert run['status'] == 'completed'
    detail = Service(service.store.root).search_detail(search['id'])
    assert len(detail['opportunities']) == 5
    assert detail['counts']['drafted'] == 3
    assert detail['counts']['company_leads'] == 1
    assert [x['rank'] for x in detail['opportunities']] == list(range(1, 6))
    assert all(not x['ready_to_contact'] and not x['ready_to_apply'] for x in detail['opportunities'])
    assert any(x['eligibility'] == 'blocked' for x in detail['opportunities'])
    for scope in ('apply', 'outreach'):
        with pytest.raises(AppError, match='No eligible'):
            service.export(search['id'], scope=scope)


def test_roles_at_same_company_are_distinct_and_tracking_urls_dedupe(service, prepared):
    search, opportunity = prepared
    seed = OpportunitySeed.model_validate({key: opportunity[key] for key in OpportunitySeed.model_fields})
    seed.listing_url += '?utm_source=board'
    seed.email = 'other@example.com'
    merged, created = service.upsert_seed(search['id'], seed)
    assert not created and merged['email'] == opportunity['email']
    seed.listing_url = opportunity['company_url'] + '/careers/999'
    seed.apply_url = seed.listing_url + '/apply'
    _, created = service.upsert_seed(search['id'], seed)
    assert created and len(service.store.list('opportunities', search['id'])) == 2


def test_apply_and_outreach_have_independent_requirements(service, prepared):
    search, opportunity = prepared
    row = service.search_detail(search['id'])['opportunities'][0]
    assert row['ready_to_apply'] and not row['ready_to_contact']
    row = ready(service, opportunity)
    assert row['ready_to_contact']
    service.suppress({'kind': 'email', 'value': opportunity['email']})
    row = service.search_detail(search['id'])['opportunities'][0]
    assert row['ready_to_apply'] and not row['ready_to_contact']
    assert row['status'] != 'suppressed'


def test_domain_exclusion_blocks_both_paths_even_after_approval(service, prepared):
    search, opportunity = prepared
    ready(service, opportunity)
    service.suppress({'kind': 'domain', 'value': opportunity['domain']})
    row = service.search_detail(search['id'])['opportunities'][0]
    assert not row['ready_to_apply'] and not row['ready_to_contact']
    with pytest.raises(AppError, match='excluded'):
        service.update_opportunity(row['id'], {'status': 'approved'})


def test_published_contact_can_be_reviewed_without_paid_verifier(service, prepared):
    search, opportunity = prepared
    assert ready(service, opportunity)['ready_to_contact']
    row = service.update_opportunity(opportunity['id'], {'email_status': 'mx_found', 'status': 'approved'})
    assert row['ready_to_contact']
    assert row['email_status'] == 'mx_found'  # Still not labelled mailbox-verified.
    service.update_search(search['id'], {'candidate': {'email': ''}})
    assert not service.search_detail(search['id'])['opportunities'][0]['ready_to_contact']


def test_stale_or_unconfirmed_posting_never_ready_to_apply(service, prepared):
    search, opportunity = prepared
    for changes in ({'last_checked_at': (datetime.now(UTC)-timedelta(days=8)).isoformat()}, {'source_check':'reachable'}, {'listing_status':'unverified'}, {'listing_source':'aggregator'}, {'valid_through':'2020-01-01'}):
        row = opportunity | changes
        service.store.save('opportunities', row)
        assert not service.search_detail(search['id'])['opportunities'][0]['ready_to_apply']


@pytest.mark.parametrize('patch', [{'contact_name':'Someone Else'}, {'linkedin_url':'https://example.com/person'}, {'email':'new@example.com'}])
def test_contact_edits_revoke_approval(service, prepared, patch):
    _, opportunity = prepared
    ready(service, opportunity)
    row = service.update_opportunity(opportunity['id'], patch)
    assert row['status'] != 'approved' and not row['ready_to_contact']
    if 'email' in patch:
        assert not row['email_source'] and row['email_status'] == 'unknown'


def test_shared_mailbox_export_once_but_application_outcome_requires_job_identity(service, prepared):
    search, opportunity = prepared
    ready(service, opportunity)
    seed = OpportunitySeed.model_validate({key: opportunity[key] for key in OpportunitySeed.model_fields})
    seed.listing_url += '-other'
    seed.apply_url += '-other'
    other, _ = service.upsert_seed(search['id'], seed)
    other.update({key: opportunity[key] for key in ('evidence','candidate_claims','match_components','drafts','research_status','source_check','last_checked_at','fit_score')})
    service.store.save('opportunities', other)
    ready(service, other)
    rows = list(csv.DictReader(io.StringIO(service.export(search['id'], scope='outreach').decode('utf-8-sig'))))
    assert len(rows) == 1
    result = service.import_outcomes(search['id'], 'progress.csv', f"email,status\n{opportunity['email']},applied\n".encode())
    assert result == {'updated':0,'skipped':1}
    result = service.import_outcomes(search['id'], 'progress.csv', f"opportunity_id,application_status\n{opportunity['id']},applied\n".encode())
    assert result['updated'] == 1
    assert service.store.get('opportunities', other['id'])['application_status'] == 'new'
    result = service.import_verification(search['id'], 'verifier.csv', f"email,status,provider,checked_at\n{opportunity['email']},catch_all,Fixture,{now()}\n".encode())
    assert result['updated'] == 2
    assert all(x['email_status'] == 'catch_all' for x in service.store.list('opportunities', search['id']))


def test_skipping_a_role_does_not_suppress_company_email(service, prepared):
    search, opportunity = prepared
    service.update_opportunity(opportunity['id'], {'application_status':'not_interested'})
    assert not service.store.suppressions()
    assert not service.search_detail(search['id'])['opportunities'][0]['ready_to_contact']


def test_bounce_suppresses_contact_but_preserves_application(service, prepared):
    search, opportunity = prepared
    service.update_opportunity(opportunity['id'], {'application_status':'bounced'})
    row = service.search_detail(search['id'])['opportunities'][0]
    assert service.is_email_suppressed(row) and not service.is_suppressed(row)
    assert row['ready_to_apply']


def test_prior_contact_in_another_search_blocks_repeat_email_not_new_job(service, prepared):
    search, opportunity = prepared
    service.update_opportunity(opportunity['id'], {'application_status':'contacted'})
    second = service.create_search({key:search[key] for key in ('name','candidate','preferences','settings','objective')})
    clone, _ = service.upsert_seed(second['id'], OpportunitySeed.model_validate({key:opportunity[key] for key in OpportunitySeed.model_fields}))
    clone.update({key:opportunity[key] for key in ('evidence','candidate_claims','match_components','drafts','research_status','source_check','last_checked_at','fit_score')})
    service.store.save('opportunities', clone)
    row = ready(service, clone)
    assert not row['ready_to_contact'] and row['ready_to_apply']
    assert any('another search' in x for x in row['export_blockers'])


def test_invalid_verifier_is_not_invalid_job_and_dates_are_real(service, prepared):
    search, opportunity = prepared
    with pytest.raises(AppError):
        service.update_opportunity(opportunity['id'], {'email_status':'verified'})
    with pytest.raises(AppError):
        service.update_opportunity(opportunity['id'], {'email_status':'verified','verification_source':'Fixture','verification_at':'2999-01-01'})
    result = service.import_verification(search['id'], 'v.csv', f"email,status,provider,checked_at\n{opportunity['email']},invalid,Fixture,{now()}\n".encode())
    assert result['updated'] == 1
    row = service.search_detail(search['id'])['opportunities'][0]
    assert row['ready_to_apply'] and row['email_status'] == 'suppressed'


def test_old_mailbox_verification_blocks_contact_export(service, prepared):
    _, opportunity = prepared
    row = service.update_opportunity(opportunity['id'], {'status':'approved','email_status':'verified','verification_source':'Fixture','verification_at':(datetime.now(UTC)-timedelta(days=31)).isoformat()})
    assert not row['ready_to_contact'] and 'Refresh mailbox verification' in row['export_blockers']


def test_upload_reads_full_material_dedupes_and_invalidates_approval(service, prepared):
    search, opportunity = prepared
    ready(service, opportunity)
    text = 'Original project experience with actual supplied evidence.'
    doc = service.upload_document(search['id'], '../resume.txt', text.encode())
    assert doc['name'] == 'resume.txt'
    assert service.upload_document(search['id'], 'copy.txt', text.encode())['duplicate']
    assert text in service.candidate_text(search['id'])
    row = service.store.get('opportunities', opportunity['id'])
    assert row['research_status'] == 'pending' and not row['drafts']
    service.delete_document(search['id'], doc['id'])
    assert text not in service.candidate_text(search['id'])


def test_combined_resume_budget_and_unknown_fields_are_rejected(service):
    search = service.create_search({'name':'Text', 'candidate':{'resume_text':'x'*59000}})
    with pytest.raises(AppError, match='60,000'):
        service.upload_document(search['id'], 'resume.txt', b'y'*1001)
    with pytest.raises(AppError, match='unknown field'):
        service.update_search(search['id'], {'profile':{'summary':'Invented'}})
    with pytest.raises(AppError, match='separate search'):
        service.update_search(search['id'], {'settings':{'provider':'demo'}})


def test_nested_route_updates_preserve_other_settings_and_change_provider_default(service):
    search = service.create_search({'name':'Models', 'settings':{'provider':'codex','model':'codex-special','routes':{'research':{'model':'research-special'}}}})
    search = service.update_search(search['id'], {'settings':{'routes':{'outreach':{'provider':'claude'}}}})
    assert route_for(search['settings'], 'research')['model'] == 'research-special'
    assert route_for(search['settings'], 'outreach') == {'provider':'claude','model':''}


def test_import_does_not_forge_live_ats_authority(service):
    search = service.create_search({'name':'Imported'})
    raw = b'company,job_title,job_url,company_website,source_provider,listing_status,last_checked_at\nExample,Engineer,https://jobs.example/123,https://example.com,ashby,open,2026-10-23\nBad,Engineer,,,,,\n'
    result = service.import_opportunities(search['id'], 'jobs.csv', raw)
    assert result['imported'] == 1 and result['rejected'] == 1
    row = service.search_detail(search['id'])['opportunities'][0]
    assert row['source_provider'] == 'import' and row['listing_status'] == 'unverified' and not row['last_checked_at']


def test_spreadsheet_contains_ranked_job_links_sources_and_complete_email(service, prepared):
    search, opportunity = prepared
    ready(service, opportunity)
    raw = service.export(search['id'])
    row = next(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
    assert row['rank'] == '1' and row['apply_url'] and row['email_3_subject']
    assert row['email_1_body'].endswith('Alex Example\nalex@example.invalid')
    assert row['score_breakdown'] and row['evidence'] and row['opportunity_id']
    book = load_workbook(io.BytesIO(service.export(search['id'], 'xlsx')))
    assert book.sheetnames == ['Report','Opportunities','Sources','Read me']
    assert book['Report']['B2'].value == search['name']
    assert book['Report']['B3'].value == search['objective']
    assert book['Report']['B12'].value == opportunity['title']
    assert book['Report']['H12'].hyperlink.target == opportunity['apply_url']
    assert book['Report'].auto_filter.ref == 'A11:H12'
    assert book['Opportunities'].freeze_panes == 'H2'


def test_no_resume_inference_from_objective_or_generated_profile(service):
    search = service.create_search({'name':'Goal', 'objective':'I want to become a flight software engineer.'})
    search['profile'] = {'summary':'Invented expertise in flight software.'}
    service.store.save('searches', search)
    original = service.candidate_text(search['id'])
    assert 'flight software' not in original
    assert not candidate_claims_supported({'candidate_claims':[{'text':'I write flight software.','resume_quote':'Invented expertise in flight software.'}]}, original)


def test_unbacked_score_and_candidate_claim_cannot_gain_points(service, prepared):
    search, opportunity = prepared
    baseline = assess(search, opportunity, service.candidate_text(search['id']))
    opportunity['match_components'][0]['candidate_quote'] = 'I led a team of 5000.'
    opportunity['match_components'][1]['source_url'] = 'https://invented.example'
    result = assess(search, opportunity, service.candidate_text(search['id']))
    assert result['fit_score'] < baseline['fit_score']
    assert result['score_breakdown'][0]['match'] == 'unknown'
    assert result['score_breakdown'][1]['match'] == 'unknown'
    opportunity['evidence'].append({'url':''})
    opportunity['match_components'][2]['source_url'] = ''
    assert assess(search, opportunity, service.candidate_text(search['id']))['score_breakdown'][2]['match'] == 'unknown'


def test_funding_cannot_dominate_fit_and_unknown_salary_is_not_mismatch(service, prepared):
    search, opportunity = prepared
    search['settings']['weights'] = {'role':1,'skills':1,'seniority':1,'location':1,'compensation':1,'preferences':1,'timing':100}
    opportunity.update(salary_min=None, salary_max=None, salary_source='', salary_period='unknown')
    result = assess(search, opportunity, service.candidate_text(search['id']))
    assert result['score_breakdown'][-1]['weight'] <= 5
    assert result['eligibility'] == 'review' and not result['blockers']


def test_explicit_sponsorship_and_work_mode_conflicts_are_separate_from_fit(service, prepared):
    search, opportunity = prepared
    search['candidate']['needs_sponsorship'] = True
    opportunity['sponsorship'] = 'no'
    search['preferences'].update(strict_work_mode=True,work_modes=['remote'])
    result = assess(search, opportunity, service.candidate_text(search['id']))
    assert result['eligibility'] == 'blocked' and len(result['blockers']) == 2


def test_company_identity_and_ats_region_do_not_merge_unrelated_jobs():
    a = {'kind':'company_lead','company_url':'','source_provider':'greenhouse','source_board':'alpha'}
    b = a | {'source_board':'beta'}
    assert identity_of(a) != identity_of(b)
    assert identity_of({'kind':'open_role','listing_url':'https://evillever.co/alpha/123'}).startswith('role:')
    assert identity_of({'kind':'open_role','listing_url':'https://jobs.lever.co/a/123'}) != identity_of({'kind':'open_role','listing_url':'https://jobs.eu.lever.co/a/123'})


def test_date_only_deadline_includes_today():
    assert not deadline_passed(now()[:10])
    assert deadline_passed('2020-01-01')
    assert not deadline_passed('nonsense')


def test_default_queries_include_news_and_rotate_roles(service):
    search = service.create_search({'name':'Discovery', 'preferences':{'roles':['Engineer','Analyst'],'locations':['LA','Seattle']}})
    plan = queries(search)
    assert len(plan) == 6 and len(set(plan)) == 6
    assert 'funding' in plan[1] and any('contract' in query for query in plan)


def test_per_company_context_is_bounded_but_quotes_remain_original(service, prepared):
    search, opportunity = prepared
    original = '\n\n'.join(f'Project {i}: Python and manufacturing data analysis, actual supplied experience.' for i in range(1000))
    excerpt = candidate_excerpt(search | {'candidate_evidence':original}, opportunity)
    assert not excerpt['complete'] and len(excerpt['text']) <= 18000
    assert all(block in original for block in excerpt['text'].split('\n\n'))


@pytest.mark.parametrize('url', ['javascript:alert(1)','file:///etc/passwd','https://u:p@example.com','data:text/html,hello'])
def test_nonweb_or_credentialed_urls_rejected(url):
    with pytest.raises(ValueError):
        public_url(url)


def test_outreach_model_change_preserves_research_and_only_invalidates_copy(service, prepared):
    search, opportunity = prepared
    ready(service, opportunity)
    service.update_search(search['id'], {'settings':{'routes':{'outreach':{'provider':'claude','model':'different'}}}})
    row = service.store.get('opportunities', opportunity['id'])
    assert row['research_status'] == 'complete' and row['match_components']
    assert not row['drafts'] and row['status'] == 'researched'
    service.update_search(search['id'], {'settings':{'routes':{'research':{'model':'different'}}}})
    assert service.store.get('opportunities', opportunity['id'])['research_status'] == 'pending'
